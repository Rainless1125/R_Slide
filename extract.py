"""從上傳檔案抽取純文字（PDF / DOCX / PPTX / Markdown / TXT），不保存檔案。

PPTX 的文字、表格與原生圖表數據直接從檔案讀出；投影片中的圖片（截圖）交給呼叫端傳入的
read_images 辨識（見 app.py），辨識結果依圖片在頁面上的位置插回同一頁，避免與其他內容錯置。
"""
import hashlib
import io
import os
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from typing import Callable

from docx import Document
from docx.table import Table
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.graphfrm import GraphicFrame
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
# 簡報常內嵌大量截圖，上限較大；仍低於 README 建議的反向代理上限（client_max_body_size 40m）
MAX_PPTX_BYTES = 30 * 1024 * 1024
MAX_REQUEST_BYTES = max(MAX_UPLOAD_BYTES, MAX_PPTX_BYTES)
# .docx / .pptx 是 zip 壓縮檔：解壓縮後的總大小上限，避免很小的檔案解壓出數 GB 的資料（zip bomb）把記憶體用光。
# 一般文件解壓後多在數 MB 以內，內嵌大量圖片時才會較大
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PPTX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".md", ".markdown", ".txt"}
OLD_FORMATS = {".doc": "Word 檔請先另存為 .docx", ".ppt": "PowerPoint 檔請先另存為 .pptx"}
TEXT_ENCODINGS = ("utf-8-sig", "cp950", "big5hkscs")


class ExtractError(Exception):
    """可直接顯示給使用者的抽取錯誤。"""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# 有明確頁碼格式的行（Page 3、- 3 -、3 / 10、第 3 頁）一律視為頁碼
PAGE_LABEL_RE = re.compile(r"page\s*\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?|-\s*\d{1,4}\s*-|\d{1,4}\s*/\s*\d{1,4}|第\s*\d{1,4}\s*頁", re.I)
# 只有數字的行可能是內文數據（年份、金額），需確認它跟著頁序遞增才移除
BARE_NUMBER_RE = re.compile(r"\d{1,4}")
# 頁碼與頁序的差值上限（封面、目錄不編號或以羅馬數字編號造成的位移）；
# 連續頁面結尾剛好是 2023、2024 這類遞增數據時，差值會大到不合理，不視為頁碼
MAX_PAGE_OFFSET = 100


def _page_number_offset(pages: list[list[str]]) -> int | None:
    """找出「頁首 / 頁尾數字 − 頁序」最常見的差值（容許封面不編號等固定位移）。
    至少一半的頁面（且至少 2 頁；單頁 PDF 則必須剛好是 1）符合才視為頁碼，否則回傳 None。"""
    counts: dict[int, int] = {}
    for i, lines in enumerate(pages, 1):
        ends = {lines[0].strip(), lines[-1].strip()} if lines else set()
        offsets = {int(line) - i for line in ends if BARE_NUMBER_RE.fullmatch(line)}
        offsets = {off for off in offsets if abs(off) <= MAX_PAGE_OFFSET}
        for off in offsets:
            counts[off] = counts.get(off, 0) + 1
    if not counts:
        return None
    if len(pages) == 1:
        return 0 if 0 in counts else None
    offset, hits = max(counts.items(), key=lambda kv: kv[1])
    return offset if hits >= max(2, (len(pages) + 1) // 2) else None


def _running_numbers(pages: list[list[str]], pos: int) -> set[str]:
    """每頁同一位置（第一行 / 最後一行）都出現的相同數字，例如頁首 / 頁尾的年份「2024」。
    至少 3 頁、且 60% 以上的頁面相同才算，避免把少數頁面剛好相同的數據當成頁首 / 頁尾。"""
    if len(pages) < 3:
        return set()
    counts: dict[str, int] = {}
    for lines in pages:
        line = lines[pos].strip() if lines else ""
        if BARE_NUMBER_RE.fullmatch(line):
            counts[line] = counts.get(line, 0) + 1
    need = max(3, -(-len(pages) * 3 // 5))   # ceil(60%)
    return {v for v, c in counts.items() if c >= need}


def _strip_page_numbers(pages: list[str]) -> list[str]:
    """只移除每頁頭尾的頁碼與頁首 / 頁尾行，避免誤刪內文中單獨成行的數據。"""
    split = [p.strip().split("\n") for p in pages]
    offset = _page_number_offset(split)
    running = {-1: _running_numbers(split, -1), 0: _running_numbers(split, 0)}

    def is_page_number(line: str, page: int, pos: int) -> bool:
        line = line.strip()
        if PAGE_LABEL_RE.fullmatch(line) or line in running[pos]:
            return True
        return offset is not None and bool(BARE_NUMBER_RE.fullmatch(line)) and int(line) - page == offset

    out = []
    for i, lines in enumerate(split, 1):
        if lines and is_page_number(lines[-1], i, -1):
            lines = lines[:-1]
        if lines and is_page_number(lines[0], i, 0):
            lines = lines[1:]
        out.append("\n".join(lines).strip())
    return out


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("此 PDF 有密碼保護，請先解除密碼後再上傳")
        pages = [page.extract_text() or "" for page in reader.pages]
    except ExtractError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise ExtractError("無法讀取此 PDF，檔案可能已損毀") from e
    text = "\n\n".join(p for p in _strip_page_numbers(pages) if p)
    if not text:
        raise ExtractError("此 PDF 抽不到文字，可能是掃描影像檔，目前不支援", 422)
    return text, len(pages)


def _table_to_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            value = " ".join(cell.text.split())
            # 合併儲存格在 python-docx 中會重複出現，略過相鄰的重複值
            if not cells or cells[-1] != value:
                cells.append(value)
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _check_zip_size(data: bytes, limit: int, label: str):
    """以 zip 目錄記錄的解壓縮大小檢查；zipfile 讀取每個檔案時最多只輸出記錄的大小，因此這個總和就是實際上限。"""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        total = sum(info.file_size for info in zf.infolist())
    if total > limit:
        raise ExtractError(f"{label}解壓縮後的內容過大，無法讀取", 413)


def _extract_docx(data: bytes) -> tuple[str, None]:
    try:
        _check_zip_size(data, MAX_DOCX_UNCOMPRESSED_BYTES, "Word 檔")
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise ExtractError("無法讀取此 Word 檔，請確認為 .docx 格式（舊版 .doc 請先另存新檔）") from e
    blocks = []
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            blocks.append(_table_to_text(item))
            continue
        text = item.text.strip()
        if not text:
            continue
        style = (item.style.name if item.style is not None else "") or ""
        level = re.match(r"Heading (\d)", style)
        if style == "Title":
            text = "# " + text
        elif level:
            text = "#" * min(int(level.group(1)) + 1, 6) + " " + text
        elif style.startswith("List"):
            text = "- " + text
        blocks.append(text)
    return "\n\n".join(b for b in blocks if b), None


def _extract_plain(data: bytes) -> tuple[str, None]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), None
    for enc in TEXT_ENCODINGS:
        try:
            return data.decode(enc), None
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), None


# ===================== PPTX =====================
# 太小的圖片多半是圖示或裝飾，不送辨識（像素寬 × 高）
MIN_IMAGE_AREA = 160 * 160
MIN_IMAGE_SIDE = 48
# 同一張圖片出現在這麼多頁以上時視為 Logo 或背景，不送辨識
REPEATED_IMAGE_SLIDES = 3
# 送辨識前把圖片長邊縮到這個大小以內（太大只會多耗 token）
MAX_IMAGE_SIDE = 2048
VISION_FORMATS = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}
SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
# 同一列的判斷範圍（投影片高度的比例）：上緣差距在此範圍內的元素視為同一列，再由左到右排序
ROW_TOLERANCE = 0.04
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
CHART_TYPE_NAMES = (("DOUGHNUT", "環圈圖"), ("PIE", "圓餅圖"), ("BAR", "橫條圖"), ("COLUMN", "直條圖"),
                    ("LINE", "折線圖"), ("AREA", "區域圖"), ("XY_SCATTER", "散佈圖"), ("RADAR", "雷達圖"))


@dataclass
class SlideImage:
    """投影片中的一張圖片。data 為空表示不送辨識（skip 寫明原因）。"""
    key: str
    data: bytes = b""
    mime: str = ""
    skip: str = ""


@dataclass
class PptxSlide:
    number: int
    title: str = ""
    blocks: list = field(default_factory=list)   # str 或 SlideImage，依畫面位置排序
    notes: str = ""

    @property
    def images(self) -> list[SlideImage]:
        return [b for b in self.blocks if isinstance(b, SlideImage) and b.data]

    def text(self) -> str:
        """本頁的文字內容（不含圖片），供辨識圖片時當作上下文。"""
        return "\n".join([self.title] + [b for b in self.blocks if isinstance(b, str)]).strip()


def _fmt_number(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(round(value, 6)) if isinstance(value, float) else str(value)


def _text_frame_text(frame) -> str:
    lines = []
    for p in frame.paragraphs:
        text = "".join(r.text for r in p.runs).strip() or p.text.strip()
        if text:
            lines.append(("  " * (p.level - 1) + "- " if p.level else "") + text)
    return "\n".join(lines)


def _pptx_table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            value = " ".join(cell.text.split())
            if not cells or cells[-1] != value:
                cells.append(value)
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _chart_text(chart) -> str:
    """原生圖表：直接讀出類別與每個系列的數值（PowerPoint 檔中存的實際數字）。"""
    type_name = getattr(chart.chart_type, "name", "") or ""
    kind = next((zh for key, zh in CHART_TYPE_NAMES if key in type_name), "圖表")
    title = ""
    if chart.has_title and chart.chart_title.has_text_frame:
        title = chart.chart_title.text_frame.text.strip()
    percent = False
    try:
        percent = "%" in (chart.value_axis.tick_labels.number_format or "")
    except (ValueError, AttributeError, KeyError):
        pass                                     # 圓餅圖等沒有數值軸
    series, categories = [], []
    for plot in chart.plots:
        if not categories:
            categories = [str(c) for c in plot.categories]
        for s in plot.series:
            values = list(s.values)
            if percent:
                values = [None if v is None else round(v * 100, 4) for v in values]
            series.append((s.name or f"系列{len(series) + 1}", values))
    if not series:
        return ""
    unit = "（%）" if percent else ""
    head = f"【圖表：{title}（{kind}）】" if title else f"【圖表（{kind}）】"
    rows = ["類別 | " + " | ".join(name + unit for name, _ in series)]
    for i, category in enumerate(categories or [str(i + 1) for i in range(len(series[0][1]))]):
        rows.append(category + " | " + " | ".join(_fmt_number(v[i]) if i < len(v) else "" for _, v in series))
    return head + "\n" + "\n".join(rows)


def _diagram_text(shape, slide) -> str:
    """SmartArt：python-pptx 不支援，從圖形的資料部分讀出所有文字。"""
    data = shape._element.xpath(".//a:graphicData/*[local-name()='relIds']/@*[local-name()='dm']")
    if not data:
        return ""
    try:
        blob = slide.part.related_part(data[0]).blob
    except KeyError:
        return ""
    texts = re.findall(rb"<a:t>([^<]*)</a:t>", blob)
    return "\n".join("- " + t.decode("utf-8", errors="replace").strip() for t in texts if t.strip())


def _shape_image(shape):
    """圖片或含圖片的版面配置區；取不到圖片時回傳 None。"""
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE and not shape.is_placeholder:
        return None
    try:
        return shape.image
    except (AttributeError, ValueError, KeyError):
        return None


def _prepare_image(blob: bytes) -> tuple[bytes, str, str]:
    """回傳 (送辨識的圖片, MIME, 略過原因)。太小的圖片略過；不支援或過大的格式轉成 PNG。"""
    try:
        with Image.open(io.BytesIO(blob)) as im:
            im.load()
            w, h = im.size
            if w * h < MIN_IMAGE_AREA or min(w, h) < MIN_IMAGE_SIDE:
                return b"", "", "small"
            if im.format in VISION_FORMATS and max(w, h) <= MAX_IMAGE_SIDE:
                return blob, VISION_FORMATS[im.format], ""
            im = im.convert("RGBA" if "A" in im.getbands() or im.mode == "P" else "RGB")
            im.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            out = io.BytesIO()
            im.save(out, "PNG", optimize=True)
            return out.getvalue(), "image/png", ""
    except (UnidentifiedImageError, OSError, ValueError):
        return b"", "", "format"                 # 例如 EMF / WMF 向量圖，Pillow 無法在 Linux 上轉換


def _shape_blocks(shape, slide, slide_no: int, counter: list) -> list:
    """一個元素轉成的內容區塊；群組依子元素的位置排序後展開。"""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [b for child in _sorted_shapes(list(shape.shapes), None) for b in _shape_blocks(child, slide, slide_no, counter)]
    if shape.is_placeholder and shape.placeholder_format.type in SKIP_PLACEHOLDERS:
        return []
    image = _shape_image(shape)
    if image is not None:
        counter[0] += 1
        return [SlideImage(key=f"s{slide_no}-i{counter[0]}", data=image.blob, mime=image.content_type or "")]
    if shape.has_chart:
        return [_chart_text(shape.chart)]
    if shape.has_table:
        table = _pptx_table_text(shape.table)
        return ["【表格】\n" + table] if table else []
    if isinstance(shape, GraphicFrame) and DIAGRAM_URI in shape._element.xml:
        return [_diagram_text(shape, slide)]
    if shape.has_text_frame:
        return [_text_frame_text(shape.text_frame)]
    return []


def _sorted_shapes(shapes: list, slide_height: int | None) -> list:
    """依閱讀順序排序：上緣相近的元素視為同一列，同一列由左到右。"""
    tops = [s.top or 0 for s in shapes]
    height = slide_height or (max(tops) - min(tops) if tops else 0) or 1
    tol = max(1, int(height * ROW_TOLERANCE))
    return sorted(shapes, key=lambda s: ((s.top or 0) // tol, s.left or 0))


def parse_pptx(data: bytes) -> tuple[list[PptxSlide], int]:
    """回傳 (各頁內容, 略過的隱藏頁數)。圖片先保留原始資料，由 prepare_pptx_images 篩選與轉檔。"""
    try:
        _check_zip_size(data, MAX_PPTX_UNCOMPRESSED_BYTES, "PowerPoint 檔")
        prs = Presentation(io.BytesIO(data))
    except ExtractError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise ExtractError("無法讀取此 PowerPoint 檔，請確認為 .pptx 格式（舊版 .ppt 請先另存新檔）") from e
    slides, hidden = [], 0
    for slide in prs.slides:
        if slide._element.get("show") in ("0", "false"):
            hidden += 1
            continue
        page = PptxSlide(number=len(slides) + 1)
        counter = [0]
        for shape in _sorted_shapes(list(slide.shapes), prs.slide_height):
            if (not page.title and shape.is_placeholder and shape.has_text_frame
                    and shape.placeholder_format.type in TITLE_PLACEHOLDERS):
                page.title = " ".join(shape.text_frame.text.split())
                continue
            page.blocks.extend(b for b in _shape_blocks(shape, slide, page.number, counter) if b)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            page.notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(page)
    return slides, hidden


def prepare_pptx_images(slides: list[PptxSlide]):
    """篩選要辨識的圖片：略過太小、格式不支援，以及在多頁重複出現的圖片（Logo、背景）。"""
    images = [b for s in slides for b in s.blocks if isinstance(b, SlideImage)]
    seen = Counter()
    for s in slides:
        for digest in {hashlib.sha256(b.data).hexdigest() for b in s.blocks if isinstance(b, SlideImage)}:
            seen[digest] += 1
    repeat_at = REPEATED_IMAGE_SLIDES if len(slides) >= REPEATED_IMAGE_SLIDES else len(slides) + 1
    for img in images:
        if seen[hashlib.sha256(img.data).hexdigest()] >= repeat_at:
            img.data, img.skip = b"", "repeated"
            continue
        img.data, img.mime, img.skip = _prepare_image(img.data)


def render_pptx(slides: list[PptxSlide], readings: dict[str, str]) -> str:
    """組成文字：每頁一段，圖片的辨識結果放在圖片原本的位置。readings 為空字串表示裝飾圖片，不輸出。"""
    out = []
    for s in slides:
        parts = [f"## 第 {s.number} 頁" + (f"：{s.title}" if s.title else "")]
        n = 0
        for b in s.blocks:
            if not isinstance(b, SlideImage):
                parts.append(b)
                continue
            if b.skip:
                continue                          # 太小、重複或格式不支援的圖片多半是裝飾，不另外標示
            n += 1
            if b.key not in readings:
                parts.append(f"【圖片 {n}：未辨識】")
            elif readings[b.key]:
                parts.append(f"【圖片 {n}（AI 辨識）】\n{readings[b.key]}")
        if s.notes:
            parts.append("【講者備註】\n" + s.notes)
        if s.title or len(parts) > 1:
            out.append("\n\n".join(parts))
    return "\n\n".join(out)


# read_images(slides) → (辨識結果 {圖片 key: 文字}, 提示訊息)；None 表示不辨識圖片
ImageReader = Callable[[list[PptxSlide]], tuple[dict[str, str], list[str]]]


def _extract_pptx(data: bytes, read_images: ImageReader | None) -> tuple[str, int, dict]:
    slides, hidden = parse_pptx(data)
    prepare_pptx_images(slides)
    images = [img for s in slides for img in s.images]
    readings, notices = (read_images(slides) if read_images and images else ({}, []))
    info = {"images": len(images), "images_read": sum(1 for img in images if img.key in readings), "notices": notices}
    unsupported = sum(1 for s in slides for b in s.blocks if isinstance(b, SlideImage) and b.skip == "format")
    if unsupported:
        info["notices"].append(f"{unsupported} 張圖片的格式（例如 EMF、WMF）無法讀取，已略過")
    if hidden:
        info["notices"].append(f"已略過 {hidden} 張隱藏投影片")
    return render_pptx(slides, readings), len(slides), info


def extract_text(filename: str, data: bytes, max_chars: int, read_images: ImageReader | None = None) -> dict:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in OLD_FORMATS:
        raise ExtractError(f"不支援舊版格式，{OLD_FORMATS[ext]}後再上傳", 415)
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractError("不支援的檔案格式，目前支援 PDF、DOCX、PPTX、Markdown、TXT", 415)
    limit = MAX_PPTX_BYTES if ext == ".pptx" else MAX_UPLOAD_BYTES
    if len(data) > limit:
        raise ExtractError(f"檔案過大（上限 {limit // 1024 // 1024} MB）", 413)
    if not data:
        raise ExtractError("檔案是空的")

    info = {}
    if ext == ".pdf":
        text, pages = _extract_pdf(data)
    elif ext == ".pptx":
        text, pages, info = _extract_pptx(data, read_images)
    elif ext == ".docx":
        text, pages = _extract_docx(data)
    else:
        text, pages = _extract_plain(data)

    text = _clean_text(text)
    if not text:
        raise ExtractError("檔案中沒有可用的文字內容", 422)
    total_chars = len(text)
    truncated = total_chars > max_chars
    return {
        "text": text[:max_chars] if truncated else text,
        "filename": os.path.basename(filename),
        "pages": pages,
        "chars": total_chars,
        "truncated": truncated,
        **info,
    }

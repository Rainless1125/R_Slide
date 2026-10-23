import base64
import io
import os

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

os.environ.setdefault("GEMINI_API_KEY", "test-key")  # app.py 啟動時需要至少一組 Key，測試不會實際呼叫 LLM

from app import app  # noqa: E402
from exporter import MAX_EXPORT_SLIDES, SLIDE_HEIGHT, SLIDE_WIDTH  # noqa: E402

client = TestClient(app)
# 檔案解析與匯出需要登入；未登入的情況見 test_requires_login
pytestmark = pytest.mark.usefixtures("logged_in")


def image_data_url(fmt: str = "JPEG", size=(256, 144), color=(30, 64, 175)) -> str:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, fmt)
    mime = "jpeg" if fmt == "JPEG" else "png"
    return f"data:image/{mime};base64," + base64.b64encode(buf.getvalue()).decode()


def export(slides, title="測試簡報"):
    return client.post("/api/export/pptx", json={"title": title, "slides": slides})


def test_builds_one_full_bleed_picture_per_slide():
    res = export([{"title": "封面", "image": image_data_url()},
                  {"title": "數據", "image": image_data_url("PNG")}])
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    prs = Presentation(io.BytesIO(res.content))
    assert (prs.slide_width, prs.slide_height) == (SLIDE_WIDTH, SLIDE_HEIGHT)
    assert prs.core_properties.title == "測試簡報"
    assert len(prs.slides) == 2
    for slide, title in zip(prs.slides, ["封面", "數據"]):
        (pic,) = slide.shapes
        assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        assert pic._element.nvPicPr.cNvPr.get("descr") == title
        assert slide.notes_slide.notes_text_frame.text == title


def test_missing_title_falls_back_to_page_number():
    res = export([{"image": image_data_url()}])
    prs = Presentation(io.BytesIO(res.content))
    assert prs.slides[0].notes_slide.notes_text_frame.text == "第 1 頁"


def test_rejects_empty_or_too_many_slides():
    assert export([]).status_code == 400
    one = {"title": "x", "image": image_data_url(size=(16, 9))}
    assert export([one] * (MAX_EXPORT_SLIDES + 1)).status_code == 400


def test_rejects_non_image_data_urls():
    svg = "data:image/svg+xml;base64," + base64.b64encode(b"<svg/>").decode()
    for bad in ["", "https://example.com/a.png", svg]:
        res = export([{"title": "x", "image": bad}])
        assert res.status_code == 400
        assert "第 1 頁" in res.json()["detail"]


def test_rejects_corrupted_image():
    res = export([{"title": "x", "image": "data:image/png;base64," + base64.b64encode(b"not an image").decode()}])
    assert res.status_code == 400


def test_rejects_mismatched_format():
    gif = io.BytesIO()
    Image.new("RGB", (8, 8)).save(gif, "GIF")
    res = export([{"title": "x", "image": "data:image/png;base64," + base64.b64encode(gif.getvalue()).decode()}])
    assert res.status_code == 400


def test_rejects_malformed_payload():
    assert client.post("/api/export/pptx", content=b"not json",
                       headers={"content-type": "application/json"}).status_code == 400
    assert client.post("/api/export/pptx", json=["x"]).status_code == 400
    assert client.post("/api/export/pptx", json={"slides": ["x"]}).status_code == 400


def test_non_string_title_falls_back_to_page_label():
    res = export([{"title": 123, "image": image_data_url()}])
    assert res.status_code == 200
    (pic,) = Presentation(io.BytesIO(res.content)).slides[0].shapes
    assert pic._element.nvPicPr.cNvPr.get("descr") == "第 1 頁"


def test_non_string_image_is_rejected():
    res = export([{"title": "封面", "image": 123}])
    assert res.status_code == 400
    assert "圖片格式不正確" in res.json()["detail"]


def test_decompression_bomb_is_rejected():
    buf = io.BytesIO()
    Image.new("1", (20000, 20000)).save(buf, "PNG")   # 壓縮後很小，但像素數超過 PIL 的安全上限
    res = export([{"title": "封面", "image": "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()}])
    assert res.status_code == 400
    assert "尺寸過大" in res.json()["detail"]


def test_requires_login(logged_in):
    from accounts import current_user
    app.dependency_overrides.pop(current_user, None)
    res = export([{"title": "封面", "image": image_data_url()}])
    assert res.status_code == 401

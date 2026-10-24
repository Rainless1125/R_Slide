import io
import os

import pytest
from docx import Document
from fastapi.testclient import TestClient

os.environ.setdefault("GEMINI_API_KEY", "test-key")  # app.py 啟動時需要至少一組 Key，測試不會實際呼叫 LLM

from app import MAX_CONTENT_CHARS, app  # noqa: E402
from extract import MAX_UPLOAD_BYTES  # noqa: E402

client = TestClient(app)
# 檔案解析與匯出需要登入；未登入的情況見 test_requires_login
pytestmark = pytest.mark.usefixtures("logged_in")


def make_pdf(pages: list[list[str]]) -> bytes:
    """產生最小的文字型 PDF（每頁數行 ASCII 文字）。"""
    n = len(pages)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [" + b" ".join(f"{4 + i * 2} 0 R".encode() for i in range(n))
            + f"] /Count {n} >>".encode(),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, lines in enumerate(pages):
        ops = "BT /F1 12 Tf 14 TL 72 720 Td " + " ".join(f"({l}) Tj T*" for l in lines) + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + i * 2} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objs.append(f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream".encode())
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    out.write(b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets))
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Q3 專案報告", level=1)
    doc.add_paragraph("核心功能完成 85%。")
    doc.add_paragraph("追加後端人力", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "指標", "數值"
    table.cell(1, 0).text, table.cell(1, 1).text = "完成度", "85%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(name: str, data: bytes):
    return client.post("/api/extract", files={"file": (name, data)})


def test_pdf_extracts_text_and_strips_page_numbers():
    res = upload("report.pdf", make_pdf([["Quarterly revenue grew", "85", "1"], ["Next steps: launch", "2"]]))
    assert res.status_code == 200
    body = res.json()
    assert body["pages"] == 2
    assert "Quarterly revenue grew" in body["text"] and "Next steps: launch" in body["text"]
    assert "85" in body["text"]  # 內文中單獨成行的數據要保留
    assert not body["text"].rstrip().endswith("2")  # 頁尾頁碼要移除
    assert body["truncated"] is False


def test_short_pdf_is_accepted():
    res = upload("short.pdf", make_pdf([["Hi"]]))
    assert res.status_code == 200
    assert res.json()["text"] == "Hi"


def test_scanned_pdf_without_text_is_rejected():
    res = upload("scan.pdf", make_pdf([[]]))
    assert res.status_code == 422
    assert "掃描" in res.json()["detail"]


def test_corrupted_pdf_is_rejected():
    res = upload("bad.pdf", b"%PDF-1.4 not really a pdf")
    assert res.status_code == 400


def test_docx_keeps_headings_lists_and_tables():
    res = upload("報告.docx", make_docx())
    assert res.status_code == 200
    text = res.json()["text"]
    assert "## Q3 專案報告" in text
    assert "- 追加後端人力" in text
    assert "指標 | 數值" in text and "完成度 | 85%" in text


def test_docx_zip_bomb_is_rejected(monkeypatch):
    import extract
    monkeypatch.setattr(extract, "MAX_DOCX_UNCOMPRESSED_BYTES", 1024)
    res = upload("bomb.docx", make_docx())
    assert res.status_code == 413
    assert "解壓縮後的內容過大" in res.json()["detail"]


def test_invalid_docx_is_rejected():
    res = upload("old.docx", b"not a zip file")
    assert res.status_code == 400


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "big5", "utf-16"])
def test_text_file_encodings(encoding):
    res = upload("note.txt", "第三季專案進度：完成 85%".encode(encoding))
    assert res.status_code == 200
    assert res.json()["text"] == "第三季專案進度：完成 85%"


def test_markdown_is_read_as_text():
    res = upload("notes.md", "# 標題\n\n\n\n\n- 重點一\r\n- 重點二".encode())
    assert res.status_code == 200
    assert res.json()["text"] == "# 標題\n\n- 重點一\n- 重點二"


def test_long_content_is_truncated():
    res = upload("long.txt", ("字" * (MAX_CONTENT_CHARS + 100)).encode())
    body = res.json()
    assert body["truncated"] is True
    assert len(body["text"]) == MAX_CONTENT_CHARS
    assert body["chars"] == MAX_CONTENT_CHARS + 100


def test_unsupported_extension_is_rejected():
    res = upload("data.xlsx", b"anything")
    assert res.status_code == 415


@pytest.mark.parametrize("name, hint", [("old.ppt", ".pptx"), ("old.doc", ".docx")])
def test_old_office_formats_explain_how_to_convert(name, hint):
    res = upload(name, b"anything")
    assert res.status_code == 415
    assert hint in res.json()["detail"]


def test_empty_file_is_rejected():
    assert upload("empty.txt", b"").status_code == 400
    assert upload("blank.txt", b"   \n\n  ").status_code == 422


def test_oversized_upload_is_rejected():
    res = upload("huge.txt", b"a" * (MAX_UPLOAD_BYTES + 128 * 1024))
    assert res.status_code == 413


def test_missing_file_field_is_rejected():
    res = client.post("/api/extract", data={"other": "x"})
    assert res.status_code == 400


@pytest.mark.parametrize("pages", [
    [["Revenue 2023", "2023"], ["Revenue 2024", "2024"], ["Users", "350"]],   # 遞增的年份不是頁碼
    [["A", "100"], ["B", "200"], ["C", "100"]],                               # 少數頁面剛好相同不是頁尾
    [["Total revenue", "350"]],
])
def test_pdf_keeps_trailing_numbers_that_are_not_page_numbers(pages):
    res = upload("data.pdf", make_pdf(pages))
    assert res.status_code == 200
    assert res.json()["text"] == "\n\n".join("\n".join(lines) for lines in pages)


def test_pdf_strips_running_header_and_footer_numbers():
    pages = [["2024", "Intro", "1"], ["2024", "Plan", "2"], ["2024", "Risks", "3"], ["2024", "End", "4"]]
    res = upload("annual.pdf", make_pdf(pages))
    assert res.status_code == 200
    assert res.json()["text"] == "Intro\n\nPlan\n\nRisks\n\nEnd"


def test_pdf_strips_offset_page_numbers_and_labels():
    res = upload("deck.pdf", make_pdf([["Cover"], ["Intro", "2"], ["Plan", "3"], ["Page 4 of 4", "End"]]))
    assert res.status_code == 200
    assert res.json()["text"] == "Cover\n\nIntro\n\nPlan\n\nEnd"


def test_requires_login(logged_in):
    from accounts import current_user
    app.dependency_overrides.pop(current_user, None)
    res = upload("notes.txt", "內容".encode())
    assert res.status_code == 401

import io
import os

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

os.environ.setdefault("GEMINI_API_KEY", "test-key")  # app.py 啟動時需要至少一組 Key，測試不會實際呼叫 LLM

import app as app_module  # noqa: E402
import store  # noqa: E402
from app import app  # noqa: E402

client = TestClient(app)
pytestmark = pytest.mark.usefixtures("logged_in")


def png(w: int, h: int, color=(30, 90, 200)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    buf.seek(0)
    return buf


def make_pptx(hidden_slide=False, logo_on_every_slide=False) -> bytes:
    prs = Presentation()
    blank, title_only = prs.slide_layouts[6], prs.slide_layouts[5]

    # 第 1 頁：標題、右側截圖、左側文字（文字在畫面上較低，但排在截圖前面的列中）
    s1 = prs.slides.add_slide(title_only)
    s1.shapes.title.text = "部門滿意度"
    s1.shapes.add_picture(png(800, 500), Inches(5), Inches(2), Inches(4), Inches(2.5))
    box = s1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(4), Inches(1)).text_frame
    box.text = "客服部最需改善"
    p = box.add_paragraph()
    p.text, p.level = "人力不足", 1
    s1.shapes.add_picture(png(20, 20), Inches(9), Inches(0.2), Inches(0.3), Inches(0.3))   # 小圖示：不送辨識
    s1.notes_slide.notes_text_frame.text = "先講結論"

    # 第 2 頁：原生圖表與表格
    s2 = prs.slides.add_slide(blank)
    data = CategoryChartData()
    data.categories = ["研發部", "行政部"]
    data.add_series("滿意度", (4.1, 4.0))
    chart = s2.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(0.5), Inches(5), Inches(3), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "各部門滿意度"
    table = s2.shapes.add_table(2, 2, Inches(0.5), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "指標", "數值"
    table.cell(1, 0).text, table.cell(1, 1).text = "回覆率", "92%"

    # 第 3 頁：群組中的截圖
    s3 = prs.slides.add_slide(blank)
    group = s3.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "流程說明"
    group.shapes.add_picture(png(600, 400, (200, 50, 50)), Inches(1), Inches(2), Inches(3), Inches(2))

    if hidden_slide:
        s4 = prs.slides.add_slide(blank)
        s4.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "備用資料"
        s4._element.set("show", "0")
    if logo_on_every_slide:
        for slide in prs.slides:
            slide.shapes.add_picture(png(400, 200, (0, 0, 0)), Inches(8), Inches(6.5), Inches(1.5), Inches(0.75))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def fake_vision(monkeypatch):
    """取代 call_llm：記錄每次呼叫收到的圖片數量，回傳可辨認的辨識結果。"""
    calls = []

    def fake(provider, model, system, prompt, schema, temperature, usage=None, images=None):
        calls.append({"provider": provider, "model": model, "prompt": prompt, "images": images})
        if usage is not None:
            usage.update(provider=provider, model=model or "fake", input_tokens=100, output_tokens=10, total_tokens=110)
        page = prompt.split("第 ")[1].split(" 頁")[0]
        return {"images": [{"kind": "chart", "content": f"辨識結果 P{page}-{i + 1}"} for i in range(len(images))]}

    monkeypatch.setattr(app_module, "call_llm", fake)
    monkeypatch.setattr(app_module, "openai_client", object())   # 測試環境只設定了 Gemini Key
    return calls


def upload(data: bytes, **form):
    return client.post("/api/extract", files={"file": ("deck.pptx", data)}, data=form)


def test_pptx_text_chart_table_notes_and_images_stay_on_their_slide(fake_vision):
    res = upload(make_pptx(), provider="openai", model="gpt-6-luna")
    assert res.status_code == 200
    body = res.json()
    text = body["text"]
    assert body["pages"] == 3
    assert body["images"] == 2 and body["images_read"] == 2

    p1, p2, p3 = text.split("## 第 ")[1:]
    assert p1.startswith("1 頁：部門滿意度")
    # 同一列由左到右：左側文字在右側截圖之前，辨識結果放在截圖原本的位置
    assert p1.index("客服部最需改善") < p1.index("辨識結果 P1-1")
    assert "- 人力不足" in p1 and "【講者備註】\n先講結論" in p1
    assert "【圖表：各部門滿意度（橫條圖）】" in p2
    assert "研發部 | 4.1" in p2 and "行政部 | 4" in p2
    assert "回覆率 | 92%" in p2
    assert "流程說明" in p3 and "辨識結果 P3-1" in p3

    # 每頁一次呼叫，附上該頁文字當上下文；小圖示不送辨識
    assert [len(c["images"]) for c in fake_vision] == [1, 1]
    assert all(c["provider"] == "openai" and c["model"] == "gpt-6-luna" for c in fake_vision)
    assert "客服部最需改善" in next(c["prompt"] for c in fake_vision if "第 1 頁" in c["prompt"])


def test_pptx_image_recognition_is_logged_as_one_usage(fake_vision, logged_in):
    upload(make_pptx(), provider="openai", model="gpt-6-luna")
    with store.connect() as conn:
        rows = conn.execute("SELECT action, status, total_tokens, model FROM llm_usage").fetchall()
    assert [tuple(r) for r in rows] == [("extract", "ok", 220, "gpt-6-luna")]


def test_pptx_without_images_does_not_call_ai(fake_vision):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "只有文字"
    buf = io.BytesIO()
    prs.save(buf)
    res = upload(buf.getvalue())
    assert res.status_code == 200
    assert "只有文字" in res.json()["text"]
    assert fake_vision == []


def test_repeated_logo_and_hidden_slides_are_skipped(fake_vision):
    res = upload(make_pptx(hidden_slide=True, logo_on_every_slide=True))
    body = res.json()
    assert body["images"] == 2                          # Logo 出現在每一頁，不送辨識
    assert "備用資料" not in body["text"]
    assert any("隱藏投影片" in n for n in body["notices"])


def test_image_limit_from_env(fake_vision, monkeypatch):
    monkeypatch.setenv("EXTRACT_IMAGE_MAX", "1")
    body = upload(make_pptx()).json()
    assert body["images_read"] == 1
    assert "【圖片 1：未辨識】" in body["text"]
    assert any("超過辨識上限" in n for n in body["notices"])


def test_images_can_be_disabled(fake_vision, monkeypatch):
    monkeypatch.setenv("EXTRACT_IMAGES", "false")
    body = upload(make_pptx()).json()
    assert fake_vision == []
    assert body["text"].count("未辨識") == 2


def test_failed_slide_does_not_block_other_slides(monkeypatch):
    def fake(provider, model, system, prompt, schema, temperature, usage=None, images=None):
        if "第 1 頁" in prompt:
            raise RuntimeError("model does not support images")
        return {"images": [{"kind": "table", "content": "第三頁表格"}]}

    monkeypatch.setattr(app_module, "call_llm", fake)
    body = upload(make_pptx()).json()
    assert "第三頁表格" in body["text"]
    assert "【圖片 1：未辨識】" in body["text"]
    assert any("辨識失敗" in n and "does not support images" in n for n in body["notices"])


def test_decoration_images_are_omitted(monkeypatch):
    monkeypatch.setattr(app_module, "call_llm",
                        lambda *a, **k: {"images": [{"kind": "decoration", "content": "背景"}]})
    text = upload(make_pptx()).json()["text"]
    assert "背景" not in text and "圖片" not in text


def test_rate_limited_user_still_gets_text(fake_vision, monkeypatch):
    monkeypatch.setenv("LLM_RATE_LIMIT_PER_HOUR", "1")
    upload(make_pptx())
    body = upload(make_pptx()).json()
    assert "客服部最需改善" in body["text"]
    assert any("圖片未辨識" in n for n in body["notices"])
    assert len(fake_vision) == 2                        # 第二次上傳沒有呼叫 AI


def test_invalid_pptx_is_rejected():
    res = client.post("/api/extract", files={"file": ("bad.pptx", b"not a zip")})
    assert res.status_code == 400


def test_pptx_zip_bomb_is_rejected(monkeypatch):
    import extract
    monkeypatch.setattr(extract, "MAX_PPTX_UNCOMPRESSED_BYTES", 1024)
    res = upload(make_pptx())
    assert res.status_code == 413


def test_call_llm_sends_images_to_openai(monkeypatch):
    sent = {}

    class Parsed:
        def model_dump(self):
            return {"images": []}

    class Completion:
        usage = None
        choices = [type("C", (), {"message": type("M", (), {"refusal": None, "parsed": Parsed(), "content": ""})()})()]

    class FakeClient:
        class beta:
            class chat:
                class completions:
                    @staticmethod
                    def parse(model, messages, response_format):
                        sent["messages"] = messages
                        return Completion()

    monkeypatch.setitem(app_module.MODEL_CATALOG, "openai", ["gpt-6-luna"])
    monkeypatch.setattr(app_module, "provider_client", lambda p: FakeClient())
    app_module.call_llm("openai", "gpt-6-luna", "sys", "看圖", app_module.SlideImageReadings, 0.2,
                        images=[(b"\x89PNG", "image/png")])
    content = sent["messages"][1]["content"]
    assert content[0] == {"type": "text", "text": "看圖"}
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")


def chart_pptx(add) -> bytes:
    prs = Presentation()
    add(prs.slides.add_slide(prs.slide_layouts[6]).shapes)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_percent_chart_values_are_converted(fake_vision):
    def add(shapes):
        data = CategoryChartData(number_format="0%")
        data.categories = ["北區", "南區"]
        data.add_series("占比", (0.6, 0.4))
        shapes.add_chart(XL_CHART_TYPE.PIE, 0, 0, Inches(4), Inches(3), data)
    text = upload(chart_pptx(add)).json()["text"]
    assert "類別 | 占比（%）\n北區 | 60\n南區 | 40" in text


def test_chart_unit_and_axis_title_are_kept(fake_vision):
    def add(shapes):
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("營收", (100, 120))
        data.add_series("成本", (80, 90))
        chart = shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), data).chart
        chart.value_axis.tick_labels.number_format = '0"萬"'
        chart.value_axis.tick_labels.number_format_is_linked = False
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = "金額"
    text = upload(chart_pptx(add)).json()["text"]
    assert "數值軸：金額" in text
    assert "類別 | 營收（萬） | 成本（萬）\nQ1 | 100 | 80\nQ2 | 120 | 90" in text


def test_scatter_chart_lists_x_and_y(fake_vision):
    def add(shapes):
        data = XyChartData()
        series = data.add_series("門市")
        series.add_data_point(1.5, 20)
        series.add_data_point(3, 45)
        shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, 0, 0, Inches(4), Inches(3), data)
    text = upload(chart_pptx(add)).json()["text"]
    assert "系列 | X | Y\n門市 | 1.5 | 20\n門市 | 3 | 45" in text

import io
import os

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

os.environ.setdefault("GEMINI_API_KEY", "test-key")  # app.py 啟動時需要至少一組 Key，測試不會實際呼叫 LLM

import app as app_module  # noqa: E402
from exporter import MAX_NOTES_CHARS  # noqa: E402
from test_export import export, image_data_url  # noqa: E402

client = TestClient(app_module.app)
pytestmark = pytest.mark.usefixtures("logged_in")   # 生成類 API 需要登入

SLIDE = {"page_number": 1, "layout_concept": "KPI", "slide_title": "營收", "html_content": "<div></div>"}


def capture_llm(monkeypatch, result):
    captured = {}

    def fake_llm(provider, model, system_instruction, prompt, schema, temperature, usage=None):
        captured.update(system=system_instruction, prompt=prompt, schema=schema)
        return dict(result)

    monkeypatch.setattr(app_module, "call_llm", fake_llm)
    return captured


def test_speaker_notes_is_required_in_llm_schema():
    schema = app_module.AdaptivePresentation.model_json_schema()
    assert "speaker_notes" in schema["$defs"]["AdaptiveSlide"]["required"]
    assert "speaker_notes" in app_module.AdaptiveSlide.model_json_schema()["required"]


def test_generate_asks_for_and_returns_speaker_notes(monkeypatch):
    captured = capture_llm(monkeypatch, {"presentation_title": "t",
                                         "slides": [{**SLIDE, "charts": [], "speaker_notes": "先說明營收成長的原因。"}]})
    res = client.post("/api/generate", json={"content": "素材"})
    assert res.status_code == 200
    assert res.json()["slides"][0]["speaker_notes"] == "先說明營收成長的原因。"
    assert "講者備註" in captured["system"]


def test_regenerate_accepts_old_slides_without_notes(monkeypatch):
    captured = capture_llm(monkeypatch, {**SLIDE, "charts": [], "speaker_notes": "新的講稿"})
    res = client.post("/api/regenerate-slide", json={"content": "素材", "slide_index": 0, "slides": [SLIDE]})
    assert res.status_code == 200
    assert res.json()["speaker_notes"] == "新的講稿"
    assert "原講者備註" in captured["prompt"] and "講者備註" in captured["system"]


def test_regenerate_sends_existing_notes_to_llm(monkeypatch):
    captured = capture_llm(monkeypatch, {**SLIDE, "charts": [], "speaker_notes": "x"})
    slides = [{**SLIDE, "speaker_notes": "強調第三季的轉折"}]
    res = client.post("/api/regenerate-slide", json={"content": "素材", "slide_index": 0, "slides": slides})
    assert res.status_code == 200
    assert "強調第三季的轉折" in captured["prompt"]


def notes_of(res):
    assert res.status_code == 200
    return [s.notes_slide.notes_text_frame.text for s in Presentation(io.BytesIO(res.content)).slides]


def test_pptx_writes_speaker_notes_and_keeps_line_breaks():
    res = export([{"title": "封面", "notes": "  大家好。\n今天談三件事。  ", "image": image_data_url()}])
    assert notes_of(res) == ["大家好。\n今天談三件事。"]


def test_pptx_falls_back_to_title_without_notes():
    res = export([{"title": "封面", "notes": "   ", "image": image_data_url()},
                  {"title": "數據", "notes": 123, "image": image_data_url()},
                  {"title": "結語", "image": image_data_url()}])
    assert notes_of(res) == ["封面", "數據", "結語"]


def test_pptx_truncates_overlong_notes():
    res = export([{"title": "x", "notes": "字" * (MAX_NOTES_CHARS + 100), "image": image_data_url()}])
    assert notes_of(res)[0] == "字" * MAX_NOTES_CHARS

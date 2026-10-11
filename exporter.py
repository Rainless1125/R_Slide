"""將前端渲染好的投影片圖片組成 PPTX，不保存檔案。"""
import base64
import binascii
import io
import re

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.util import Emu

MAX_EXPORT_SLIDES = 60
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_IMAGE_SIDE = 4096
# 16:9，與 PowerPoint「寬螢幕」預設尺寸相同（13.333 x 7.5 in）
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)
DATA_URL_RE = re.compile(r"data:image/(png|jpeg);base64,([A-Za-z0-9+/=\s]+)")


class ExportError(Exception):
    """可直接顯示給使用者的匯出錯誤。"""


def _decode_image(data_url: str, page: int) -> bytes:
    m = DATA_URL_RE.fullmatch(data_url or "")
    if not m:
        raise ExportError(f"第 {page} 頁的圖片格式不正確")
    try:
        raw = base64.b64decode(m.group(2), validate=False)
    except (binascii.Error, ValueError) as e:
        raise ExportError(f"第 {page} 頁的圖片資料損毀") from e
    if len(raw) > MAX_IMAGE_BYTES:
        raise ExportError(f"第 {page} 頁的圖片過大")
    try:
        with Image.open(io.BytesIO(raw)) as img:
            if img.format not in ("PNG", "JPEG"):
                raise ExportError(f"第 {page} 頁的圖片格式不正確")
            if max(img.size) > MAX_IMAGE_SIDE:
                raise ExportError(f"第 {page} 頁的圖片尺寸過大")
            img.verify()
    except (UnidentifiedImageError, OSError, SyntaxError) as e:
        raise ExportError(f"第 {page} 頁的圖片資料損毀") from e
    return raw


def build_pptx(title: str, slides: list[dict]) -> bytes:
    """每頁放一張滿版圖片；投影片標題寫入圖片替代文字與備忘稿。"""
    if not slides:
        raise ExportError("沒有可匯出的投影片")
    if len(slides) > MAX_EXPORT_SLIDES:
        raise ExportError(f"投影片超過 {MAX_EXPORT_SLIDES} 頁，無法匯出")
    images = [_decode_image(s.get("image", ""), i) for i, s in enumerate(slides, 1)]

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
    prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    for i, (slide_data, raw) in enumerate(zip(slides, images), 1):
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(io.BytesIO(raw), 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        slide_title = (slide_data.get("title") or "").strip() or f"第 {i} 頁"
        pic.name = f"Slide {i}"
        pic._element.nvPicPr.cNvPr.set("descr", slide_title)
        slide.notes_slide.notes_text_frame.text = slide_title

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
